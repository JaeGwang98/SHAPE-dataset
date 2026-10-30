# -*- coding: utf-8 -*-
"""R7-1, R8-1, R9-3: RI 증강 샘플 의사 평가용 시트 (질병당 5개 x 26 = 130, 층화)."""
import json, pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
D = "../2025_05_19_data"
rows = []
for src, fs, tot in [("GPT-4o-mini", ["batch_4o-mini_50000_random_fewshot.jsonl", "batch_4o-mini_28000_random_fewshot.jsonl"], "4o_mini_total.csv"),
                     ("GPT-4.1-mini", ["batch_4.1-mini_50000_random_fewshot.jsonl", "batch_4.1-mini_28000_random_fewshot.jsonl"], "4.1_mini_total.csv")]:
    P = []
    for f in fs: P += [json.loads(l)["body"]["messages"][1]["content"] for l in open(f"{D}/{f}", encoding="utf-8")]
    t = pd.read_csv(f"{D}/{tot}")
    cur = [p.split("현재증상:")[1] for p in P]
    for i in range(len(P)):
        if i % 5:
            rows.append({"source": src, "row": i, "variant": i % 5, "질문": t["질문"][i], "질병명": t["labels"][i],
                         "original_symptoms": cur[i - i % 5], "inserted_symptom": cur[i][len(cur[i - i % 5]) + 2:]})
df = pd.DataFrame(rows)
rng = np.random.default_rng(2026)
pick = []
for lab, g in df.groupby("질병명"):
    # 질병당 5개: 모델 교대, 서로 다른 원본 행
    g = g.sample(frac=1, random_state=int(rng.integers(1e9)))
    g = g.drop_duplicates(subset=["source", "original_symptoms"])
    a = g[g.source == "GPT-4o-mini"].head(3); b = g[g.source == "GPT-4.1-mini"].head(2)
    pick.append(pd.concat([a, b]))
S = pd.concat(pick).sample(frac=1, random_state=7).reset_index(drop=True)
S.insert(0, "연번", range(1, len(S) + 1))
key = S[["연번", "질병명", "source", "row", "variant", "original_symptoms", "inserted_symptom"]]
blank = S[["연번", "질문", "질병명"]].copy()
for c in ["① 실제 환자 질문과 유사한가 (1-4)", "② 질병명이 질문 내용에 적절한가 (예/아니오)", "③ 증상 조합이 임상적으로 개연성 있는가 (1-4)", "④ (②가 아니오인 경우) 더 적절한 질병명", "비고"]:
    blank[c] = ""
guide = pd.DataFrame({"평가 안내": [
    "아래 130개 문항은 SHAPE 학습 데이터 중 '무작위 증상 삽입(RI)'이 적용된 샘플입니다 (질병당 5개, 26개 질병).",
    "각 문항은 원래 증상 목록에 해당 질병과 무관할 수 있는 증상 1개가 추가된 상태에서 생성되었습니다. 어떤 증상이 추가되었는지는 공개하지 않습니다(맹검).",
    "① 실제 환자 질문과 유사한가: 1=전혀 아니다, 2=다소 아니다, 3=다소 그렇다, 4=매우 그렇다",
    "② 질병명 적절성: 제시된 질병명이 질문에 기술된 증상 전체를 고려할 때 여전히 적절한 진단 후보인지 (예/아니오)",
    "③ 증상 조합의 임상적 개연성: 1=불가능/모순, 2=드묾, 3=가능, 4=흔함",
    "④ ②에서 '아니오'라면 더 적절하다고 생각하는 질병명을 26개 목록 중에서 적어 주세요(복수 가능).",
    "평가는 다른 평가자와 상의 없이 독립적으로 진행해 주세요.",
    "26개 질병 목록: " + ", ".join(sorted(df["질병명"].unique()))]})
out = "round2/RI_expert_review_sheet.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as w:
    guide.to_excel(w, "안내", index=False)
    for r in ["평가자A", "평가자B", "평가자C"]: blank.to_excel(w, r, index=False)
    key.to_excel(w, "KEY_평가자에게_보내지_말것", index=False)
wb = load_workbook(out)
for ws in wb.worksheets:
    for col in ws.columns:
        L = col[0].column_letter; ws.column_dimensions[L].width = 90 if col[0].value in ("질문", "평가 안내", "original_symptoms") else 18
        for c in col: c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
wb.save(out)
print(len(S), S["질병명"].value_counts().unique(), S["source"].value_counts().to_dict(), S["variant"].value_counts().to_dict())
print(key.head(5).to_string())
